#!/usr/bin/env python3
"""Generate data/master.xlsx from data/mapping.yaml + data/languages/*.yaml.

The mapping is data (data/mapping.yaml — the reviewable artifact); this
script only validates and renders. Validation: every non-X phoneme of every
language must be mapped, and no two phonemes of one language may land on the
same form unless data/mapping.yaml lists them under `merges`.

Rows: every input form in use, grouped by keystroke count. Columns: keys,
input, type, then the form's readings from the most to the least used, each
as IPA / languages / native speakers (millions; native, because totals
double-count multilinguals). X (marginal) phonemes are excluded.
"""
import glob
import os
import sys
import unicodedata as ud

import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

C_POINTS = {'sj','zj','tq','dq','tc','dc','kx','gx','kq','hq','rc','tx','dx','nx','py','ky','gy'}
RETRO = {'ty','dy','ny','ly','ry'}
SPECIAL2 = {'tcy','tqy','dqy','tyy','dyy','ryy'}
V_POINTS = {'ij','ej','uj','oj','aj','ic','ec','uc','oc','ac','eq','iq','aq'}
SEQ = {'aix','aux','ojix','eix','oux','ojicx','ejix','eqicx','ajux'}


def form_type(f):
    if len(f) == 1: return 'base'
    if f in C_POINTS or f in V_POINTS: return 'point'
    if f.endswith('yx') or f in SPECIAL2: return 'special'
    if f in RETRO: return 'retro'
    if f in SEQ: return 'sequence'
    if f in {'ix','ux','icx'}: return 'glide'
    if f.endswith('w'): return 'long'
    if f.endswith('y'): return 'nasal'
    if f.endswith('x'): return 'shift'
    if f.endswith('j'): return 'soft'
    return '?'


def nfd(s):
    return ud.normalize('NFD', s)


def load_mapping():
    m = yaml.safe_load(open(os.path.join(ROOT, 'data', 'mapping.yaml')))
    cons = {nfd(k): v for k, v in m['consonants'].items()}
    vows = {nfd(k): v for k, v in m['vowels'].items()}
    over = {code: {nfd(k): v for k, v in d.items()} for code, d in m['overrides'].items()}
    merges = {code: [{nfd(a), nfd(b)} for a, b in pairs]
              for code, pairs in (m.get('merges') or {}).items()}
    return cons, vows, over, merges


def aggregate():
    cons, vows, over, merges = load_mapping()
    agg, errors = {}, []
    for path in sorted(glob.glob(os.path.join(ROOT, 'data', 'languages', '*.yaml'))):
        d = yaml.safe_load(open(path))
        code, native = d['code'], d['speakers']['native']
        taken = {}
        for sec, table in (('consonants', cons), ('vowels', vows)):
            for ipa, freq in d[sec].items():
                if freq == 'X':
                    continue
                key = nfd(ipa)
                form = over.get(code, {}).get(key) or table.get(key)
                if not form:
                    errors.append(f'{code}: unmapped {ipa}')
                    continue
                if form in taken:
                    pair = {key, nfd(taken[form])}
                    if pair not in merges.get(code, []):
                        errors.append(f'{code}: {ipa} and {taken[form]} both on {form}')
                        continue
                taken[form] = ipa
                e = agg.setdefault(form, {}).setdefault(ud.normalize('NFC', ipa), [0, 0.0])
                e[0] += 1
                e[1] += native
    if errors:
        print('MAPPING INVALID:', *errors, sep='\n  ')
        sys.exit(1)
    return agg


def main():
    agg = aggregate()
    rows = []
    for form, reads in agg.items():
        rd = sorted(reads.items(), key=lambda kv: -kv[1][1])
        total = sum(s for _, (_, s) in reads.items())
        rows.append((len(form), -total, form, form_type(form), rd))
    rows.sort()
    max_r = max(len(r[4]) for r in rows)

    wb = Workbook()
    ws = wb.active
    ws.title = 'master'
    ipa_font = Font(name='Arial', size=11)
    head_font = Font(name='Arial', size=11, bold=True)
    group_fill = PatternFill('solid', fgColor='EEEEEE')
    header = ['keys', 'input', 'type']
    for i in range(1, max_r + 1):
        header += [f'ipa {i}', f'langs {i}', f'native M {i}']
    ws.append(header)
    for c in ws[1]:
        c.font = head_font
    prev_keys = 0
    for keys, _, form, t, rd in rows:
        row = [keys, form, t]
        for ipa, (n, spk) in rd:
            row += [ipa, n, round(spk)]
        ws.append(row)
        for c in ws[ws.max_row]:
            c.font = ipa_font
            if keys != prev_keys:
                c.fill = group_fill
        prev_keys = keys
    widths = {'A': 6, 'B': 9, 'C': 10}
    for i in range(max_r):
        for off, w in ((0, 8), (1, 7), (2, 10)):
            col = ws.cell(row=1, column=4 + i * 3 + off).column_letter
            widths[col] = w
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'
    out = os.path.join(ROOT, 'data', 'master.xlsx')
    wb.save(out)
    n_reads = sum(len(r[4]) for r in rows)
    print(f'mapping valid · {out}: {len(rows)} forms, {n_reads} readings, max {max_r} per form')


if __name__ == '__main__':
    main()
